# SPDX-FileCopyrightText: 2026 Pace Wisdom Solutions Pvt. Ltd.
# SPDX-License-Identifier: Apache-2.0

import io
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from curriculum.models import (
    Module, Node, Task, TaskSubmission,
    Quiz, QuizSubmission, StudentNodeProgress
)

FONT_FAMILY = "Segoe UI"
STATUS_NEEDS_MANUAL_REVIEW = "Needs Manual Review"
HEADER_AWARDED_SCORE_RESULT = "Awarded Score / Result"
HEADER_SUBMISSION_ATTEMPT_DETAILS = "Submission / Attempt Details"

def _get_learning_material_activity(node, status, last_accessed_date):
    mat = node.learning_material
    content_url_val = mat.content_url
    if not content_url_val:
        content_url_val = mat.content_file.url if mat.content_file else 'Rich Text'
    
    sub_status_val = 'Locked'
    if status == 'Completed':
        sub_status_val = 'Read'
    elif status == 'In_Progress':
        sub_status_val = 'Not Read'

    return {
        'type': 'Learning Material',
        'title': f"Material: {mat.content_type}",
        'sub_status': sub_status_val,
        'score': 'N/A',
        'details': f"Type: {mat.content_type} | URL: {content_url_val}",
        'feedback': 'N/A',
        'date': last_accessed_date if status == 'Completed' else None
    }

def _get_task_activity(node, student):
    task = node.task
    latest_sub = TaskSubmission.objects.filter(task=task, student=student, is_deleted=False).order_by('-submitted_at').first()
    if latest_sub:
        total_attempts = TaskSubmission.objects.filter(task=task, student=student, is_deleted=False).count()
        attempt_number = TaskSubmission.objects.filter(task=task, student=student, submitted_at__lte=latest_sub.submitted_at, is_deleted=False).count()
        
        details_parts = []
        if latest_sub.payload:
            details_parts.append(f"Payload: {latest_sub.payload}")
        if latest_sub.submission_file:
            details_parts.append(f"File: {latest_sub.submission_file.name}")
        details_str = " | ".join(details_parts) or "No details"
        
        return {
            'type': 'Task',
            'title': task.title,
            'sub_status': latest_sub.status,
            'score': str(latest_sub.awarded_score) if latest_sub.awarded_score is not None else 'Pending',
            'details': f"Attempt #{attempt_number} of {total_attempts} | {details_str}",
            'feedback': latest_sub.feedback or '-',
            'date': latest_sub.submitted_at
        }
    
    return {
        'type': 'Task',
        'title': task.title,
        'sub_status': 'Not Submitted',
        'score': 'N/A',
        'details': '-',
        'feedback': '-',
        'date': None
    }

def _get_quiz_activities(student, quiz_list):
    activities = []
    for q in quiz_list:
        latest_sub = QuizSubmission.objects.filter(quiz=q, student=student, is_deleted=False).order_by('-submitted_at').first()
        if latest_sub:
            total_attempts = QuizSubmission.objects.filter(quiz=q, student=student, is_deleted=False).count()
            activities.append({
                'type': 'Quiz',
                'title': q.name,
                'sub_status': latest_sub.status,
                'score': f"{latest_sub.score}%",
                'details': f"Attempt #{latest_sub.attempt_number} of {total_attempts} | {latest_sub.correct_answers}/{latest_sub.total_questions} correct",
                'feedback': 'Auto-Graded',
                'date': latest_sub.submitted_at
            })
        else:
            activities.append({
                'type': 'Quiz',
                'title': q.name,
                'sub_status': 'Not Attempted',
                'score': 'N/A',
                'details': '-',
                'feedback': '-',
                'date': None
            })
    return activities

def _get_node_activities_list(node, student, status, last_accessed_date):
    activities = []
    
    # 1. Learning Material
    if hasattr(node, 'learning_material') and node.learning_material and not node.learning_material.is_deleted:
        activities.append(_get_learning_material_activity(node, status, last_accessed_date))
        
    # 2. Task
    if hasattr(node, 'task') and node.task and not node.task.is_deleted:
        activities.append(_get_task_activity(node, student))
        
    # 3. Quizzes
    quizzes = Quiz.objects.filter(node=node, is_deleted=False).order_by('id')
    if quizzes.exists():
        activities.extend(_get_quiz_activities(student, quizzes))
        
    if not activities:
        activities.append({
            'type': 'Empty Node',
            'title': '-',
            'sub_status': 'N/A',
            'score': 'N/A',
            'details': '-',
            'feedback': '-',
            'date': None
        })
        
    return activities

def _write_report_headers(ws, course, student, completion_percentage, completed_nodes_count, total_nodes_count, styles):
    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = "STUDENT COURSE PROGRESS REPORT"
    title_cell.font = styles['title_font']
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40
    
    export_date_str = timezone.now().strftime('%Y-%m-%d %H:%M:%S UTC')
    metadata = [
        ("Student Name:", f"{student.first_name} {student.last_name}".strip() or student.email, "Email:", student.email),
        ("Course Title:", course.title, "Progress:", f"{round(completion_percentage, 2)}% ({completed_nodes_count}/{total_nodes_count} completed)"),
        ("Export Date:", export_date_str, "Organization:", course.organization.name)
    ]
    
    for idx, (lbl1, val1, lbl2, val2) in enumerate(metadata):
        row_num = 3 + idx
        ws.row_dimensions[row_num].height = 20
        
        ws[f"A{row_num}"] = lbl1
        ws[f"A{row_num}"].font = styles['section_font']
        ws[f"B{row_num}"] = val1
        ws[f"B{row_num}"].font = styles['normal_font']
        
        ws[f"D{row_num}"] = lbl2
        ws[f"D{row_num}"].font = styles['section_font']
        ws[f"E{row_num}"] = val2
        ws[f"E{row_num}"].font = styles['normal_font']
        
    headers = [
        "Level (Module)",
        "Chapter (Node)",
        "Chapter Status",
        "Activity Type",
        "Activity Title",
        "Submission Status",
        HEADER_AWARDED_SCORE_RESULT,
        HEADER_SUBMISSION_ATTEMPT_DETAILS,
        "Teacher Feedback",
        "Submitted Date (UTC)"
    ]
    
    ws.row_dimensions[7].height = 28
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=7, column=col_idx)
        cell.value = header
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = styles['thin_border']

def _write_activity_row(ws, row, module_title, node_title, status, act, styles):
    ws.row_dimensions[row].height = 20
    
    normal_font = styles['normal_font']
    thin_border = styles['thin_border']
    
    ws.cell(row=row, column=1, value=module_title).font = normal_font
    ws.cell(row=row, column=1).border = thin_border
    
    ws.cell(row=row, column=2, value=node_title).font = normal_font
    ws.cell(row=row, column=2).border = thin_border
    
    # Chapter status
    status_cell = ws.cell(row=row, column=3, value=status.replace('_', ' '))
    status_cell.border = thin_border
    status_cell.alignment = Alignment(horizontal="center", vertical="center")
    if status in styles['status_fills']:
        status_cell.fill = styles['status_fills'][status]
        status_cell.font = styles['status_fonts'][status]
    else:
        status_cell.font = normal_font
        
    ws.cell(row=row, column=4, value=act['type']).font = normal_font
    ws.cell(row=row, column=4).border = thin_border
    
    ws.cell(row=row, column=5, value=act['title']).font = normal_font
    ws.cell(row=row, column=5).border = thin_border
    
    # Submission status
    sub_status_cell = ws.cell(row=row, column=6, value=act['sub_status'])
    sub_status_cell.border = thin_border
    sub_status_cell.alignment = Alignment(horizontal="center", vertical="center")
    sub_status_raw = act['sub_status']
    if sub_status_raw in styles['submission_fills']:
        sub_status_cell.fill = styles['submission_fills'][sub_status_raw]
        sub_status_cell.font = styles['submission_fonts'][sub_status_raw]
    else:
        sub_status_cell.font = normal_font
        
    ws.cell(row=row, column=7, value=act['score']).font = normal_font
    ws.cell(row=row, column=7).border = thin_border
    ws.cell(row=row, column=7).alignment = Alignment(horizontal="center", vertical="center")
    
    ws.cell(row=row, column=8, value=act['details']).font = normal_font
    ws.cell(row=row, column=8).border = thin_border
    ws.cell(row=row, column=8).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    ws.cell(row=row, column=9, value=act['feedback']).font = normal_font
    ws.cell(row=row, column=9).border = thin_border
    ws.cell(row=row, column=9).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    date_val = act['date'].strftime('%Y-%m-%d %H:%M') if act['date'] else '-'
    ws.cell(row=row, column=10, value=date_val).font = normal_font
    ws.cell(row=row, column=10).border = thin_border
    ws.cell(row=row, column=10).alignment = Alignment(horizontal="center", vertical="center")

def _auto_adjust_column_widths(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 7:
                continue
            val_str = str(cell.value or '')
            if '\n' in val_str:
                val_str = max(val_str.split('\n'), key=len)
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

def _get_excel_styles(font_family=FONT_FAMILY):
    return {
        'title_font': Font(name=font_family, size=16, bold=True, color="1F4E78"),
        'section_font': Font(name=font_family, size=11, bold=True, color="595959"),
        'normal_font': Font(name=font_family, size=11),
        'header_font': Font(name=font_family, size=11, bold=True, color="FFFFFF"),
        'header_fill': PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid"),
        'thin_border': Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        ),
        'status_fills': {
            'Completed': PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
            'In_Progress': PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
            'Unlocked': PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"),
            'Locked': PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"),
        },
        'status_fonts': {
            'Completed': Font(name=font_family, size=11, color="375623", bold=True),
            'In_Progress': Font(name=font_family, size=11, color="7F6000", bold=True),
            'Unlocked': Font(name=font_family, size=11, color="1F4E78", bold=True),
            'Locked': Font(name=font_family, size=11, color="7F7F7F"),
        },
        'submission_fills': {
            'Approved': PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
            'Passed': PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
            'Accepted': PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
            'Rejected': PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
            'Failed': PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
            'Pending': PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
        },
        'submission_fonts': {
            'Approved': Font(name=font_family, size=11, color="375623", bold=True),
            'Passed': Font(name=font_family, size=11, color="375623", bold=True),
            'Accepted': Font(name=font_family, size=11, color="375623", bold=True),
            'Rejected': Font(name=font_family, size=11, color="C65911", bold=True),
            'Failed': Font(name=font_family, size=11, color="C65911", bold=True),
            'Pending': Font(name=font_family, size=11, color="7F6000", bold=True),
        }
    }

def _calculate_progress(course, student):
    from django.db.models import Q
    nodes_with_content = Node.objects.filter(module__course=course, is_deleted=False).filter(
        Q(learning_material__isnull=False) | 
        Q(task__isnull=False) | 
        Q(assessment__isnull=False) | 
        Q(quizzes__isnull=False)
    ).distinct()
    
    total_nodes_count = nodes_with_content.count()
    completed_nodes_count = StudentNodeProgress.objects.filter(
        student=student,
        node__in=nodes_with_content,
        status='Completed',
        node__is_deleted=False
    ).values('node_id').distinct().count()
    completion_percentage = (completed_nodes_count / total_nodes_count * 100) if total_nodes_count > 0 else 0
    return completion_percentage, completed_nodes_count, total_nodes_count

def _process_node_activities(ws, m_title, n, student, current_row, styles):
    prog = StudentNodeProgress.objects.filter(student=student, node=n, is_deleted=False).first()
    status = prog.status if prog else 'Locked'
    last_accessed_date = prog.last_accessed if prog else None
    
    activities = _get_node_activities_list(n, student, status, last_accessed_date)
    # Skip empty nodes (nodes that only return the 'Empty Node' placeholder activity)
    if len(activities) == 1 and activities[0]['type'] == 'Empty Node':
        return current_row
        
    for act in activities:
        _write_activity_row(ws, current_row, m_title, n.title, status, act, styles)
        current_row += 1
    return current_row

def _write_modules_progress(ws, course, student, start_row, styles):
    current_row = start_row
    modules = Module.objects.filter(course=course, is_deleted=False).order_by('sequence_order')
    
    for m in modules:
        nodes = Node.objects.filter(module=m, is_deleted=False).order_by('sequence_order')
        for n in nodes:
            current_row = _process_node_activities(ws, m.title, n, student, current_row, styles)
    return current_row

def _populate_student_detailed_sheet(ws, course, student, styles):
    ws.sheet_view.showGridLines = True
    
    # Calculate progress percentage
    completion_percentage, completed_nodes_count, total_nodes_count = _calculate_progress(course, student)
    
    # Write Title, Metadata, Headers
    _write_report_headers(ws, course, student, completion_percentage, completed_nodes_count, total_nodes_count, styles)
        
    _write_modules_progress(ws, course, student, start_row=8, styles=styles)

    _auto_adjust_column_widths(ws)

def generate_progress_excel(course, student):
    wb = Workbook()
    ws = wb.active
    ws.title = "Detailed Progress"
    
    styles = _get_excel_styles()
    _populate_student_detailed_sheet(ws, course, student, styles)
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def clean_sheet_title(title):
    invalid_chars = [':', '\\', '/', '?', '*', '[', ']']
    for char in invalid_chars:
        title = title.replace(char, '')
    return title.strip()[:30]

def _get_quiz_submission_cell_data(quiz, student_user, student_id):
    from django.db.models import Q
    subs = QuizSubmission.objects.filter(quiz=quiz, is_deleted=False)
    if student_user:
        subs = subs.filter(
            Q(student=student_user) | Q(student__user=student_user) | Q(student__uuid=student_id)
        )
    else:
        subs = subs.filter(
            Q(student_id=student_id) | Q(student__user_id=student_id) | Q(student__uuid=student_id)
        )
    latest_sub = subs.order_by('-submitted_at').first()
    if latest_sub:
        total_attempts = subs.count()
        score_str = f"{latest_sub.score}%"
        details_str = f"Attempt #{latest_sub.attempt_number} of {total_attempts} | {latest_sub.correct_answers}/{latest_sub.total_questions} correct"
    else:
        score_str = "N/A"
        details_str = "-"
    return quiz.name, score_str, details_str

def _write_quiz_summary_cells(ws, row, quizzes, max_quizzes, student_user, student_id, styles):
    for q_idx in range(max_quizzes):
        start_col = 7 + (q_idx * 3)
        if q_idx < len(quizzes):
            quiz_name, score_str, details_str = _get_quiz_submission_cell_data(
                quizzes[q_idx], student_user, student_id
            )
        else:
            quiz_name, score_str, details_str = "-", "-", "-"

        ws.cell(row=row, column=start_col, value=quiz_name).font = styles['normal_font']
        ws.cell(row=row, column=start_col).border = styles['thin_border']

        cell_score = ws.cell(row=row, column=start_col + 1, value=score_str)
        cell_score.font = styles['normal_font']
        cell_score.border = styles['thin_border']
        cell_score.alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(row=row, column=start_col + 2, value=details_str).font = styles['normal_font']
        ws.cell(row=row, column=start_col + 2).border = styles['thin_border']

def _populate_summary_row(ws, row, item, course_quizzes_map, max_quizzes, styles):
    from accounts.models import User
    ws.row_dimensions[row].height = 20

    ws.cell(row=row, column=1, value=item['learner_name']).font = styles['normal_font']
    ws.cell(row=row, column=1).border = styles['thin_border']

    ws.cell(row=row, column=2, value=item.get('learner_email', '')).font = styles['normal_font']
    ws.cell(row=row, column=2).border = styles['thin_border']

    ws.cell(row=row, column=3, value=item['course_title']).font = styles['normal_font']
    ws.cell(row=row, column=3).border = styles['thin_border']

    completion_cell = ws.cell(row=row, column=4, value=f"{item['completion_percentage']}%")
    completion_cell.font = styles['normal_font']
    completion_cell.border = styles['thin_border']
    completion_cell.alignment = Alignment(horizontal="center", vertical="center")

    modules_cell = ws.cell(row=row, column=5, value=item['modules_progress'])
    modules_cell.font = styles['normal_font']
    modules_cell.border = styles['thin_border']
    modules_cell.alignment = Alignment(horizontal="center", vertical="center")

    last_act = item.get('last_activity')
    if isinstance(last_act, timezone.datetime):
        date_str = last_act.strftime('%Y-%m-%d %H:%M')
    elif last_act:
        date_str = str(last_act)
    else:
        date_str = '-'

    date_cell = ws.cell(row=row, column=6, value=date_str)
    date_cell.font = styles['normal_font']
    date_cell.border = styles['thin_border']
    date_cell.alignment = Alignment(horizontal="center", vertical="center")

    student_id = item.get('student_id')
    course_id = item.get('course_id')
    quizzes = course_quizzes_map.get(course_id, [])
    student_user = User.objects.filter(id=student_id).first()

    _write_quiz_summary_cells(ws, row, quizzes, max_quizzes, student_user, student_id, styles)

def _create_student_sheets(wb, results, styles):
    from accounts.models import User
    from curriculum.models import Course
    used_titles = set()
    for item in results:
        student_id = item['student_id']
        course_id = item['course_id']

        try:
            student = User.objects.get(id=student_id)
            course = Course.objects.get(id=course_id)
        except (User.DoesNotExist, Course.DoesNotExist):
            continue

        base_title = f"{student.first_name} {student.last_name}".strip() or student.email
        sheet_title = clean_sheet_title(base_title)

        unique_title = sheet_title
        counter = 2
        while unique_title.lower() in used_titles:
            suffix = f" ({counter})"
            max_len = 31 - len(suffix)
            unique_title = clean_sheet_title(base_title[:max_len]) + suffix
            counter += 1

        used_titles.add(unique_title.lower())
        ws_student = wb.create_sheet(title=unique_title)
        _populate_student_detailed_sheet(ws_student, course, student, styles)

def generate_bulk_progress_excel(results, org_name):
    wb = Workbook()
    
    # 1. Summary Sheet: Overall Progress
    ws_summary = wb.active
    ws_summary.title = "Overall Progress"
    ws_summary.sheet_view.showGridLines = True
    
    styles = _get_excel_styles()
    font_family = FONT_FAMILY
    
    # Headers
    headers = ["Learner", "Email ID", "Course", "Completion %", "Modules Progress", "Last Activity"]

    # Pre-fetch quizzes per course
    course_ids = {item['course_id'] for item in results if 'course_id' in item}
    course_quizzes_map = {
        c_id: list(Quiz.objects.filter(
            node__module__course_id=c_id,
            is_deleted=False,
            node__is_deleted=False,
            node__module__is_deleted=False
        ).order_by('node__module__sequence_order', 'node__sequence_order', 'id'))
        for c_id in course_ids
    }

    max_quizzes = max((len(q_list) for q_list in course_quizzes_map.values()), default=0)

    if max_quizzes == 1:
        headers.extend(["Quiz Name", HEADER_AWARDED_SCORE_RESULT, HEADER_SUBMISSION_ATTEMPT_DETAILS])
    elif max_quizzes > 1:
        for i in range(1, max_quizzes + 1):
            headers.extend([f"Quiz {i} Name", HEADER_AWARDED_SCORE_RESULT, HEADER_SUBMISSION_ATTEMPT_DETAILS])

    # Write Title spanning all columns
    last_col_letter = get_column_letter(len(headers))
    ws_summary.merge_cells(f"A1:{last_col_letter}1")
    title_cell = ws_summary["A1"]
    title_cell.value = "LEARNER OVERALL PROGRESS REPORT"
    title_cell.font = Font(name=font_family, size=16, bold=True, color="1F4E78")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.row_dimensions[1].height = 40
    
    # Metadata
    export_date_str = timezone.now().strftime('%Y-%m-%d %H:%M:%S UTC')
    metadata = [
        ("Organization:", org_name, "Export Date:", export_date_str),
    ]
    
    for idx, (lbl1, val1, lbl2, val2) in enumerate(metadata):
        row_num = 3 + idx
        ws_summary.row_dimensions[row_num].height = 20
        ws_summary[f"A{row_num}"] = lbl1
        ws_summary[f"A{row_num}"].font = styles['section_font']
        ws_summary[f"B{row_num}"] = val1
        ws_summary[f"B{row_num}"].font = styles['normal_font']
        ws_summary[f"D{row_num}"] = lbl2
        ws_summary[f"D{row_num}"].font = styles['section_font']
        ws_summary[f"E{row_num}"] = val2
        ws_summary[f"E{row_num}"].font = styles['normal_font']
        
    ws_summary.row_dimensions[5].height = 28
    for col_idx, header in enumerate(headers, start=1):
        cell = ws_summary.cell(row=5, column=col_idx)
        cell.value = header
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = styles['thin_border']
        
    # Populate Summary Rows
    for idx, item in enumerate(results, start=6):
        _populate_summary_row(ws_summary, idx, item, course_quizzes_map, max_quizzes, styles)
        
    _auto_adjust_column_widths(ws_summary)
    _create_student_sheets(wb, results, styles)
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
