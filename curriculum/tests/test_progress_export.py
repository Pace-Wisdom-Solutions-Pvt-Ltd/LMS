# SPDX-FileCopyrightText: 2026 Pace Wisdom Solutions Pvt. Ltd.
# SPDX-License-Identifier: Apache-2.0

import pytest
from io import BytesIO
from openpyxl import load_workbook
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APIClient

from curriculum.models import (
    Course, Module, Node, StudentNodeProgress, Task, TaskSubmission,
    Quiz, QuizSubmission
)
from organizations.models import Organization, OrganizationMember
from rbac.models import Role
from accounts.models import User

@pytest.fixture
def api_client():
    return APIClient()

@pytest.fixture
def setup_data(db):
    org = Organization.objects.create(name="Test Org", slug="test-org")
    
    # Roles
    admin_role, _ = Role.objects.get_or_create(name="org_admin")
    teacher_role, _ = Role.objects.get_or_create(name="teacher")
    student_role, _ = Role.objects.get_or_create(name="student")
    
    # Users
    admin_user = User.objects.create_user(email="admin@test.com", password="pw", first_name="Admin", last_name="User")
    teacher_user = User.objects.create_user(email="teacher@test.com", password="pw", first_name="Teacher", last_name="User")
    student_user = User.objects.create_user(email="student@test.com", password="pw", first_name="Student", last_name="User")
    other_student = User.objects.create_user(email="other@test.com", password="pw", first_name="Other", last_name="User")
    
    # Organization Membership
    OrganizationMember.objects.create(user=admin_user, organization=org, role=admin_role)
    OrganizationMember.objects.create(user=teacher_user, organization=org, role=teacher_role)
    OrganizationMember.objects.create(user=student_user, organization=org, role=student_role)
    OrganizationMember.objects.create(user=other_student, organization=org, role=student_role)
    
    # Course Structure
    course = Course.objects.create(organization=org, title="Test Course", status="Published")
    module = Module.objects.create(course=course, title="Module 1: Getting Started", sequence_order=1)
    
    # Nodes
    node1 = Node.objects.create(module=module, title="Node 1: Introduction", sequence_order=1)
    node2 = Node.objects.create(module=module, title="Node 2: Final Quiz", sequence_order=2)
    
    # Task under Node 1
    task = Task.objects.create(node=node1, title="Intro Task")
    TaskSubmission.objects.create(
        task=task,
        student=student_user,
        payload="https://mycode.com",
        status="Approved",
        awarded_score=90,
        feedback="Excellent job!"
    )
    
    # Quiz under Node 2
    quiz = Quiz.objects.create(node=node2, name="Final Quiz MCQ")
    QuizSubmission.objects.create(
        quiz=quiz,
        student=student_user,
        score=75.0,
        total_questions=4,
        correct_answers=3,
        passed=True,
        status="Passed",
        attempt_number=1
    )

    # Progress (created AFTER task/quiz to avoid deletion via signals)
    StudentNodeProgress.objects.create(student=student_user, node=node1, status="Completed")
    StudentNodeProgress.objects.create(student=student_user, node=node2, status="In_Progress")
    
    return {
        "org": org,
        "admin": admin_user,
        "teacher": teacher_user,
        "student": student_user,
        "other_student": other_student,
        "course": course,
        "node1": node1,
        "node2": node2,
    }


@pytest.mark.django_db
def test_export_all_progress(api_client, setup_data):
    from organizations.models import Batch, BatchStudent
    
    org = setup_data["org"]
    course = setup_data["course"]
    student = setup_data["student"]
    teacher = setup_data["teacher"]
    
    # 1. Create a Batch
    batch = Batch.objects.create(organization=org, name="Test Batch", start_date="2026-01-01", end_date="2026-12-31")
    batch.courses.add(course)
    
    # 2. Enroll student in the Batch
    BatchStudent.objects.create(batch=batch, student=student)
    
    # 3. Associate course with teacher
    course.teachers.add(teacher)
    
    # Test bulk export URL
    url = reverse("export-all-learners-progress", kwargs={"org_id": org.id})
    
    # 4. Anonymous access is blocked
    resp = api_client.get(url, {"course_id": course.id, "teacher_id": str(teacher.id)})
    assert resp.status_code == status.HTTP_401_UNAUTHORIZED
    
    # 5. Student cannot access
    api_client.force_authenticate(user=student)
    resp = api_client.get(url, {"course_id": course.id, "teacher_id": str(teacher.id)})
    assert resp.status_code == status.HTTP_403_FORBIDDEN
    
    # 6. Teacher can access
    api_client.force_authenticate(user=teacher)
    resp = api_client.get(url, {"course_id": course.id, "teacher_id": str(teacher.id)})
    assert resp.status_code == status.HTTP_200_OK
    assert resp['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert "attachment; filename=" in resp['Content-Disposition']
    
    # Load and verify Excel sheet tabs
    wb = load_workbook(filename=BytesIO(resp.content))
    sheet_names = wb.sheetnames
    assert "Overall Progress" in sheet_names
    assert "Student User" in sheet_names
    
    # Verify overall progress summary
    ws_summary = wb["Overall Progress"]
    assert ws_summary["A1"].value == "LEARNER OVERALL PROGRESS REPORT"
    assert ws_summary["A6"].value == "Student User"
    assert ws_summary["B6"].value == "student@test.com"
    assert ws_summary["C6"].value == "Test Course"
    # Verify quiz columns beside Last Activity
    assert ws_summary["G5"].value == "Quiz Name"
    assert ws_summary["H5"].value == "Awarded Score / Result"
    assert ws_summary["I5"].value == "Submission / Attempt Details"
    assert ws_summary["G6"].value == "Final Quiz MCQ"
    assert ws_summary["H6"].value == "75.00%"
    assert ws_summary["I6"].value == "Attempt #1 of 1 | 3/4 correct"
    
    # Verify detailed sheet for the student
    ws_student = wb["Student User"]
    assert ws_student["A1"].value == "STUDENT COURSE PROGRESS REPORT"
    
    # 7. Test missing teacher_id parameter
    resp = api_client.get(url, {"course_id": course.id})
    assert resp.status_code == status.HTTP_400_BAD_REQUEST


@pytest.mark.django_db
def test_export_all_progress_duplicate_names(api_client, setup_data):
    from organizations.models import Batch, BatchStudent
    
    org = setup_data["org"]
    course = setup_data["course"]
    student = setup_data["student"]  # Student User
    teacher = setup_data["teacher"]
    
    # Create another student with the same first_name and last_name
    dup_student = User.objects.create_user(
        email="dup@test.com", password="pw",
        first_name=student.first_name, last_name=student.last_name
    )
    from rbac.models import Role
    from organizations.models import OrganizationMember
    student_role = Role.objects.get(name="student")
    OrganizationMember.objects.create(user=dup_student, organization=org, role=student_role)
    
    # Enroll both in batch
    batch = Batch.objects.create(organization=org, name="Test Batch", start_date="2026-01-01", end_date="2026-12-31")
    batch.courses.add(course)
    BatchStudent.objects.create(batch=batch, student=student)
    BatchStudent.objects.create(batch=batch, student=dup_student)
    
    course.teachers.add(teacher)
    
    url = reverse("export-all-learners-progress", kwargs={"org_id": org.id})
    api_client.force_authenticate(user=teacher)
    resp = api_client.get(url, {"course_id": course.id, "teacher_id": str(teacher.id)})
    assert resp.status_code == status.HTTP_200_OK
    
    wb = load_workbook(filename=BytesIO(resp.content))
    sheet_names = wb.sheetnames
    assert "Overall Progress" in sheet_names
    assert "Student User" in sheet_names
    assert "Student User (2)" in sheet_names


@pytest.mark.django_db
def test_export_all_progress_multiple_quizzes_dynamic_columns(api_client, setup_data):
    from organizations.models import Batch, BatchStudent

    org = setup_data["org"]
    course = setup_data["course"]
    student = setup_data["student"]
    teacher = setup_data["teacher"]
    node2 = setup_data["node2"]

    # Add a second quiz to node2
    quiz2 = Quiz.objects.create(node=node2, name="Advanced Quiz 2")

    # Enroll student in batch
    batch = Batch.objects.create(organization=org, name="Batch Multi Quiz", start_date="2026-01-01", end_date="2026-12-31")
    batch.courses.add(course)
    BatchStudent.objects.create(batch=batch, student=student)
    course.teachers.add(teacher)

    url = reverse("export-all-learners-progress", kwargs={"org_id": org.id})
    api_client.force_authenticate(user=teacher)
    resp = api_client.get(url, {"course_id": course.id, "teacher_id": str(teacher.id)})
    assert resp.status_code == status.HTTP_200_OK

    wb = load_workbook(filename=BytesIO(resp.content))
    ws_summary = wb["Overall Progress"]

    # Quiz 1 columns
    assert ws_summary["G5"].value == "Quiz 1 Name"
    assert ws_summary["H5"].value == "Awarded Score / Result"
    assert ws_summary["I5"].value == "Submission / Attempt Details"
    assert ws_summary["G6"].value == "Final Quiz MCQ"
    assert ws_summary["H6"].value == "75.00%"
    assert ws_summary["I6"].value == "Attempt #1 of 1 | 3/4 correct"

    # Quiz 2 columns (unattempted by student)
    assert ws_summary["J5"].value == "Quiz 2 Name"
    assert ws_summary["K5"].value == "Awarded Score / Result"
    assert ws_summary["L5"].value == "Submission / Attempt Details"
    assert ws_summary["J6"].value == "Advanced Quiz 2"
    assert ws_summary["K6"].value == "N/A"
    assert ws_summary["L6"].value == "-"


